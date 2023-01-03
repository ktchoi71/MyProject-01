from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart

wb = load_workbook("sample3.xlsx")
ws = wb.active

bar_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
bar_chart = BarChart()
bar_chart.add_data(bar_value)

lien_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(lien_value, titles_from_data=True)
line_chart.title = "성적표"
line_chart.style = 20
line_chart.y_axis.title = "점수"
line_chart.x_axis.title = "번호"


ws.add_chart(bar_chart, "E1")
ws.add_chart(line_chart, "N1")

wb.save("sample3_chart.xlsx")

